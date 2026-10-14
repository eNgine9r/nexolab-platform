#!/usr/bin/env python3
"""Isolated native ARM64 Energy export benchmark. Never point at a site database.

Run only inside a disposable, memory/CPU/PID-capped candidate container connected
to a dedicated non-published PostgreSQL test network. All records are synthetic.
No Modbus, hardware or production APIs are contacted.
"""
from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from math import isfinite
import json
import os
import threading
import time
import uuid
import zipfile
import re
import urllib.request
import urllib.error
from datetime import UTC, datetime, timedelta
from pathlib import Path

from sqlalchemy import func, insert, select
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session

from app.contracts import TelemetryEvent
from app.db import Database, TelemetryLatest, TelemetrySample
from app.energy_export.jobs import ExportJobs, ExportPolicy
from app.energy_export.schemas import METRICS, ExportOptions, SeriesSelection
from app.ingestion import TelemetryIngestor
from app.model_registry import register_models
from app.nodes.models import CentralNode
from app.security.repository import SecurityRepository
from app.state import RuntimeState

ORG = "13641364-1364-1364-1364-136413641364"
NODE = "energy-1364-synthetic"
OWNER = (ORG, "isolated-benchmark", "synthetic-reader", "synthetic-session")
LIMITS = {"service_memory": 512 * 1024**2, "export_rss": 256 * 1024**2,
          "spool_per_job": 512 * 1024**2, "spool_total": 2 * 1024**3,
          "disk_reserve": 128 * 1024**2, "job_seconds": 600}


def rss(pid: int | None) -> int:
    if not pid:
        return 0
    try:
        for line in Path(f"/proc/{pid}/status").read_text().splitlines():
            if line.startswith("VmRSS:"):
                return int(line.split()[1]) * 1024
    except (FileNotFoundError, ProcessLookupError, PermissionError):
        pass
    return 0


def cgroup_bytes() -> int:
    try:
        return int(Path("/sys/fs/cgroup/memory.current").read_text())
    except (FileNotFoundError, ValueError):
        return -1


def make_series() -> list[SeriesSelection]:
    result = []
    for meter in ("LE01MP-200", "LE01MP-201", "LE01MP-202", "LE01MP-203", "SDM120M-1"):
        for metric, (_, unit, _) in METRICS.items():
            if meter.startswith("SDM120M") and metric == "temperature.internal":
                continue
            result.append(SeriesSelection(node_id=NODE, equipment_id=meter,
                                          channel_id=f"{meter}-{metric}", metric=metric, unit=unit))
    assert len(result) == 44, len(result)
    return result


def seed(database: Database, series: list[SeriesSelection], *, ticks: int, base: datetime) -> int:
    security = SecurityRepository(database)
    security.provision_organization(organization_id=ORG, slug="energy-1364-synthetic", name="Energy 1364 synthetic")
    with Session(database.engine) as session, session.begin():
        session.add(CentralNode(id=str(uuid.uuid4()), organization_id=ORG, node_id=NODE,
                                display_name="Synthetic Energy 1364", state="active",
                                created_by="isolated-benchmark", clock_status="ok",
                                clock_warning_ms=30000, clock_critical_ms=120000))
    total = 0
    block = []
    latest = []
    for series_index, series_item in enumerate(series):
        final_event = None
        for tick in range(ticks):
            if series_index == 0 and tick == 12:
                continue  # One genuinely missing observation; never fill it.
            event_id = str(uuid.UUID(int=100000000 + series_index * (ticks + 2) + tick))
            is_counter = series_item.metric == "electrical.energy.active"
            value = (tick * 0.007 + series_index) if is_counter else (220 + series_index * 0.125 + tick * 0.001)
            if is_counter and tick >= 1000:
                value = (tick - 1000) * 0.007
            quality = "valid"
            if tick == 11:
                quality, value = "sensor_error", None
            row = {**series_item.model_dump(), "event_id": event_id,
                   "captured_at": base + timedelta(seconds=tick * 180),
                   "received_at": datetime.now(UTC), "value": value, "quality": quality,
                   "source": "energy-1364-synthetic", "raw_payload": {}}
            block.append(row)
            final_event = row
            total += 1
            if len(block) >= 1000:
                with database.engine.begin() as conn:
                    conn.execute(insert(TelemetrySample), block)
                block.clear()
        assert final_event is not None
        latest.append({k: final_event[k] for k in ("event_id", "node_id", "captured_at", "metric",
                        "value", "unit", "quality", "source", "equipment_id", "channel_id", "received_at") }
                      )
    # Same timestamp, a different original event identity must remain present.
    tie = dict(block[-1] if block else {**series[0].model_dump()})
    tie.update({**series[0].model_dump(), "event_id": str(uuid.uuid4()),
                "captured_at": base + timedelta(seconds=180), "received_at": datetime.now(UTC),
                "value": 219.125, "quality": "valid", "source": "energy-1364-synthetic-tie", "raw_payload": {}})
    block.append(tie)
    total += 1
    with database.engine.begin() as conn:
        if block:
            conn.execute(insert(TelemetrySample), block)
        # Synthetic source must have a genuine latest projection so an actual
        # application startup does not trigger the 10,000-row gap guard.
        # Never weaken the production reconciliation bound for a test fixture.
        for row in latest:
            row["sample_id"] = int(conn.execute(
                select(TelemetrySample.id).where(TelemetrySample.event_id == row["event_id"])
            ).scalar_one())
        conn.execute(insert(TelemetryLatest), latest)
    return total


def until(jobs: ExportJobs, job_id: str, target: str, deadline: float, samples: dict) -> dict:
    while time.monotonic() < deadline:
        value = jobs.get(OWNER, job_id)
        proc = jobs._process
        pid = proc.pid if proc is not None else None
        samples["rss_peak"] = max(samples["rss_peak"], rss(pid))
        memory = cgroup_bytes()
        if memory >= 0:
            samples["cgroup_peak"] = max(samples["cgroup_peak"], memory)
        if value["status"] == target:
            return value
        if value["status"] in {"failed", "canceled", "expired"}:
            raise RuntimeError(f"Export failed: {value['error']}")
        time.sleep(0.1)
    raise TimeoutError(f"Export did not reach {target} within fixed time budget")



def verify_existing_xlsx(database_url: str, root: Path) -> None:
    """Independent saved-source to RAW XLSX reconciliation, offline/streamed."""
    from openpyxl import load_workbook
    database = Database(database_url)
    try:
        with database.engine.connect() as conn:
            reference = {row[0]: row[1] for row in conn.execute(
                select(TelemetrySample.event_id, TelemetrySample.value)
                .where(TelemetrySample.source.like("energy-1364-synthetic%"))
            )}
        if len(reference) < 100000:
            raise ValueError("Expected >100000 original records")
        files = []
        for manifest in root.glob("*/manifest.json"):
            entry = json.loads(manifest.read_text(encoding="utf-8"))
            options = entry.get("options", {})
            if (entry.get("status") == "ready" and options.get("format") == "xlsx"
                    and options.get("step_seconds") == 0):
                files.append(manifest.parent / "artifact.xlsx")
        if len(files) != 1:
            raise ValueError(f"Expected one ready RAW XLSX in private snapshot, found {len(files)}")
        seen = set()
        workbook = load_workbook(files[0], read_only=True, data_only=True)
        rows = 0
        try:
            for sheet in workbook.worksheets:
                if not sheet.title.startswith("Data"):
                    continue
                data = sheet.iter_rows(values_only=True)
                headings = next(data)
                checks = [(i, headings.index(title.removesuffix(" · event_id")))
                          for i, title in enumerate(headings) if isinstance(title, str)
                          and title.endswith(" · event_id")]
                if not checks:
                    raise AssertionError("XLSX raw event identity columns missing")
                for row in data:
                    rows += 1
                    for identity_index, number_index in checks:
                        identity = row[identity_index] if identity_index < len(row) else None
                        if not identity or identity not in reference:
                            continue
                        observed = row[number_index]
                        expected = reference[identity]
                        assert observed is None if expected is None else (
                            isinstance(observed, (int, float)) and abs(observed-expected) < 1e-6
                        ), (identity, observed, expected)
                        if identity in seen:
                            raise AssertionError(f"Original event duplicated in XLSX: {identity}")
                        seen.add(identity)
        finally:
            workbook.close()
        if seen != reference.keys():
            raise AssertionError(f"Missing source identity: XLSX {len(seen)} / PostgreSQL {len(reference)}")
        print(json.dumps({"result": "PASS", "kind": "original_postgresql_to_raw_xlsx",
                          "source_records": len(reference), "verified_identity_values": len(seen),
                          "xlsx_rows": rows, "artifact_bytes": files[0].stat().st_size,
                          "classification": "native-isolated-synthetic-not-production"},
                         sort_keys=True), flush=True)
    finally:
        database.dispose()




def verify_existing_aggregates(database_url: str, root: Path) -> None:
    """Cross-check 10-minute source aggregates against independent PostgreSQL and XLSX."""
    from openpyxl import load_workbook
    from app.energy_export.schemas import normalized_unit

    entries = []
    for manifest in root.glob("*/manifest.json"):
        value = json.loads(manifest.read_text(encoding="utf-8"))
        if (value["status"] == "ready"
                and value["options"].get("step_seconds") == 600):
            entries.append((value["options"]["format"], manifest.parent, value))
    assert len(entries) == 2, f"Expected one 10min CSV and XLSX: {len(entries)}"
    by_format = {item[0]: item for item in entries}
    assert set(by_format) == {"csv", "xlsx"}
    _, csv_dir, csv_info = by_format["csv"]
    _, xlsx_dir, xlsx_info = by_format["xlsx"]
    csv_path = csv_dir / "artifact.csv"
    xlsx_path = xlsx_dir / "artifact.xlsx"
    opts = csv_info["options"]
    start = datetime.fromisoformat(opts["from"].replace("Z", "+00:00"))
    end = datetime.fromisoformat(opts["to"].replace("Z", "+00:00"))
    snapshot = datetime.fromisoformat(csv_info["preview"]["snapshot_at"].replace("Z", "+00:00"))
    upper_id = int(csv_info["preview"]["source_id_watermark"])
    expected = defaultdict(list)
    database = Database(database_url)
    try:
        with Session(database.engine) as session:
            result = session.execute(select(
                TelemetrySample.equipment_id, TelemetrySample.metric,
                TelemetrySample.captured_at, TelemetrySample.value,
                TelemetrySample.quality, TelemetrySample.unit
            ).where(
                TelemetrySample.node_id == NODE,
                TelemetrySample.id <= upper_id,
                TelemetrySample.received_at <= snapshot,
                TelemetrySample.captured_at >= start,
                TelemetrySample.captured_at < end,
            ).order_by(TelemetrySample.id))
            for equipment, metric, captured, value, quality, unit in result:
                if metric not in METRICS:
                    continue
                if metric == "temperature.internal" and equipment == "SDM120M-1":
                    continue
                bucket = int((captured - start).total_seconds() // 600)
                usable = (quality == "valid" and value is not None and isfinite(value)
                          and normalized_unit(unit) == normalized_unit(METRICS[metric][1])
                          and (metric != "electrical.energy.active" or value >= 0))
                expected[equipment, bucket, metric].append((captured, value if usable else None))
    finally:
        database.dispose()

    verified_cells = 0
    verified_buckets = 0
    counter_resets = 0
    mapping = {}
    seen = set()
    with csv_path.open(encoding="utf-8", newline="") as source:
        reader = csv.DictReader(source)
        csv_headers = list(reader.fieldnames or [])
        for metric, (symbol, unit, _) in METRICS.items():
            normalized_label = "°C" if unit == "degC" else "PF (ratio)" if unit == "ratio" else unit
            mapping[metric] = f"{symbol}, {normalized_label}"
        for row in reader:
            equipment = row["equipment_id"]
            bucket = int((datetime.fromisoformat(row["bucket_start_utc"].replace("Z", "+00:00")) - start).total_seconds() // 600)
            key = (row["bucket_start_utc"], equipment)
            assert key not in seen, f"duplicate bucket/meter {key}"
            seen.add(key)
            verified_buckets += 1
            for metric, header in mapping.items():
                if header + " · source count" not in row:
                    if equipment == "SDM120M-1" and metric == "temperature.internal":
                        continue
                    raise AssertionError(f"Missing expected metric {metric}")
                readings = expected[equipment, bucket, metric]
                good = [float(value) for _, value in readings if value is not None]
                bad = len(readings) - len(good)
                for suffix, expected_count in (
                        ("source count", len(readings)),
                        ("valid count", len(good)),
                        ("invalid count", bad)):
                    actual = int(row[header + " · " + suffix] or 0)
                    assert actual == expected_count, (equipment, bucket, metric, suffix, actual, expected_count)
                    verified_cells += 1
                if metric == "electrical.energy.active":
                    if "counter_decrease" in row[header + " · delta status"]:
                        counter_resets += 1
                        assert row[header + " · delta kWh"] == "", "Counter decrease must not generate energy delta"
                    continue
                for suffix, calc in (("avg", (sum(good) / len(good)) if good else None),
                                     ("min", min(good) if good else None),
                                     ("max", max(good) if good else None)):
                    cell = row[header + " · " + suffix]
                    if calc is None:
                        assert cell == "", (equipment, bucket, metric, suffix, cell)
                    else:
                        assert cell != "" and abs(float(cell) - calc) < 1e-6, (equipment, bucket, metric, suffix, cell, calc)
                    verified_cells += 1
    assert verified_buckets == csv_info["artifact"]["output_rows"]
    assert counter_resets >= 4, f"Reset not independently detected for expected counters: {counter_resets}"

    # The XLSX job has a DIFFERENT immutable snapshot. Never compare it to the
    # CSV snapshot; independently recalculate every XLSX source row/metric.
    expected.clear()
    xopts = xlsx_info["options"]
    xstart = datetime.fromisoformat(xopts["from"].replace("Z", "+00:00"))
    xend = datetime.fromisoformat(xopts["to"].replace("Z", "+00:00"))
    xsnapshot = datetime.fromisoformat(xlsx_info["preview"]["snapshot_at"].replace("Z", "+00:00"))
    xupper_id = int(xlsx_info["preview"]["source_id_watermark"])
    xexpected = defaultdict(list)
    database = Database(database_url)
    try:
        with Session(database.engine) as session:
            source = session.execute(select(
                TelemetrySample.equipment_id, TelemetrySample.metric,
                TelemetrySample.captured_at, TelemetrySample.value,
                TelemetrySample.quality, TelemetrySample.unit
            ).where(
                TelemetrySample.node_id == NODE,
                TelemetrySample.id <= xupper_id,
                TelemetrySample.received_at <= xsnapshot,
                TelemetrySample.captured_at >= xstart,
                TelemetrySample.captured_at < xend,
            ).order_by(TelemetrySample.id))
            for equipment, metric, captured, value, quality, unit in source:
                if metric not in METRICS:
                    continue
                if metric == "temperature.internal" and equipment == "SDM120M-1":
                    continue
                bucket = int((captured - xstart).total_seconds() // 600)
                usable = (quality == "valid" and value is not None and isfinite(value)
                          and normalized_unit(unit) == normalized_unit(METRICS[metric][1])
                          and (metric != "electrical.energy.active" or value >= 0))
                xexpected[equipment, bucket, metric].append(value if usable else None)
    finally:
        database.dispose()
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    xlsx_rows = 0
    typed_numeric = 0
    verified_xlsx_cells = 0
    xlsx_reset_invalidation = 0
    seen_xlsx = set()
    try:
        for sheet in workbook.worksheets:
            if not sheet.title.startswith("Data"):
                continue
            records = sheet.iter_rows(values_only=True)
            headings = next(records)
            idx = {name: i for i, name in enumerate(headings)}
            assert "bucket_start_utc" in idx and "equipment_id" in idx
            for record in records:
                xlsx_rows += 1
                equipment = record[idx["equipment_id"]]
                bucket = int((datetime.fromisoformat(record[idx["bucket_start_utc"]].replace("Z", "+00:00")) - xstart).total_seconds() // 600)
                key = (equipment,bucket)
                assert key not in seen_xlsx, f"duplicate XLSX meter/bucket {key}"
                seen_xlsx.add(key)
                for metric, header in mapping.items():
                    if metric == "temperature.internal" and equipment == "SDM120M-1":
                        continue
                    readings = xexpected[equipment, bucket, metric]
                    valid = [float(v) for v in readings if v is not None]
                    for suffix,count in (
                        ("source count",len(readings)),("valid count",len(valid)),
                        ("invalid count",len(readings)-len(valid))):
                        actual = record[idx[header + " · " + suffix]]
                        assert int(actual or 0) == count, (equipment,bucket,metric,suffix,actual,count)
                        verified_xlsx_cells += 1
                    if metric == "electrical.energy.active":
                        if "counter_decrease" in str(record[idx[header + " · delta status"]]):
                            xlsx_reset_invalidation += 1
                            assert record[idx[header + " · delta kWh"]] is None
                        continue
                    for suffix, calculated in (("avg",sum(valid)/len(valid) if valid else None),
                                               ("min",min(valid) if valid else None),
                                               ("max",max(valid) if valid else None)):
                        actual = record[idx[header + " · " + suffix]]
                        if calculated is None:
                            assert actual is None, (equipment,bucket,metric,suffix,actual)
                        else:
                            assert isinstance(actual,(int,float)) and abs(actual-calculated)<1e-6, (
                                equipment,bucket,metric,suffix,actual,calculated)
                            typed_numeric += 1
                        verified_xlsx_cells += 1
    finally:
        workbook.close()
    assert xlsx_rows == xlsx_info["artifact"]["output_rows"]
    assert typed_numeric > 1000
    assert xlsx_reset_invalidation >= 4
    print(json.dumps({"result":"PASS", "kind":"independent_aggregate_postgresql_and_xlsx",
        "csv_buckets":verified_buckets, "csv_verified_source_counts_and_statistics":verified_cells,
        "csv_counter_reset_invalidation_buckets":counter_resets,
        "xlsx_rows":xlsx_rows,"xlsx_verified_source_counts_and_statistics":verified_xlsx_cells,
        "xlsx_counter_reset_invalidation_buckets":xlsx_reset_invalidation,
        "typed_numeric_xlsx_cells":typed_numeric,
        "snapshot_boundaries":"independent_per_artifact",
        "classification":"native-isolated-synthetic-not-production"},
        ensure_ascii=False,sort_keys=True),flush=True)


def api_benchmark(url: str) -> None:
    """HTTP export and 10/s normal synthetic ingestion with fixed baseline gates."""
    endpoint = "http://nexolab-energy-1364-api-synth-v3:8082"
    database = Database(url)
    source_label = f"energy-1364-api-{uuid.uuid4().hex[:12]}"
    state = RuntimeState()
    ingestor = TelemetryIngestor(database, state, queue_maxsize=500)
    from datetime import timezone
    series = make_series()[0]
    accept = {}
    elapsed = {}
    errors = []
    samples = []
    records = []
    stop = threading.Event()
    phase = {"value": "before"}
    start = time.monotonic()
    def percentile(v, q):
        order = sorted(v)
        return order[min(len(order)-1, max(0, int(len(order)*q)-1))] if order else float("inf")
    def request(path, *, method="GET", body=None, timeout=15):
        payload = None if body is None else json.dumps(body, separators=(",",":")).encode()
        rq = urllib.request.Request(
            endpoint+path, data=payload, method=method,
            headers={"Content-Type": "application/json"} if payload is not None else {})
        started = time.monotonic()
        try:
            with urllib.request.urlopen(rq, timeout=timeout) as response:
                content = response.read()
                return response.status, json.loads(content) if content else None, round((time.monotonic()-started)*1000,4)
        except urllib.error.HTTPError as exc:
            raise RuntimeError(f"isolated API status={exc.code}, path={path}, body={exc.read()[:500]!r}") from exc
    def monitor():
        while not stop.is_set():
            try:
                code, data, latency = request("/health/ready", timeout=2)
                if code!=200 or data.get("status")!="ready":
                    errors.append(("http_ready",code))
                samples.append((phase["value"],latency))
            except Exception as exc:
                errors.append(("http_error",str(exc)[:90]))
            stop.wait(.2)
    def feeder():
        index=0
        while not stop.is_set():
            event = TelemetryEvent(event_id=uuid.uuid4(), **series.model_dump(),
                captured_at=datetime.now(UTC), value=220.0+index*.0001,
                quality="valid", source=source_label)
            when = time.monotonic()
            good = ingestor.submit_payload(json.dumps(event.normalized_payload()).encode())
            accept[str(event.event_id)]=(when, phase["value"], good)
            index += 1
            stop.wait(.1)
    ingestor.start()
    polling=threading.Thread(target=monitor,daemon=True)
    feeding=threading.Thread(target=feeder,daemon=True)
    try:
        status,catalog,latency=request("/api/v1/energy/exports/catalog")
        if status!=200 or len(catalog["items"])!=44:
            raise AssertionError(f"Expected 44 catalog series, got {len(catalog['items'])}")
        polling.start();feeding.start()
        time.sleep(30) # 30-second live-ingest *before-load* baseline.
        baseline=[ms for label,ms in samples if label=="before"]
        if len(baseline)<100:
            raise AssertionError(f"Insufficient API baseline samples: {len(baseline)}")
        phase["value"]="during"
        for format in ("csv","xlsx"):
            now=datetime.now(UTC)
            options={"series":[{k:item[k] for k in ("node_id","equipment_id","channel_id","metric","unit")}
                               for item in catalog["items"]],
                "from":(now-timedelta(days=7)).isoformat(),
                "to":now.isoformat(),"step_seconds":0,"format":format}
            begin=time.monotonic()
            code,job,_=request("/api/v1/energy/exports/preview",method="POST",body=options)
            assert code==202
            jid=job["job_id"]
            def until_status(expected):
                deadline=begin+600
                while time.monotonic()<deadline:
                    code, data, latency = request(f"/api/v1/energy/exports/{jid}")
                    assert code==200
                    if data["status"]==expected: return data
                    if data["status"] in ("failed","expired","canceled"):
                        raise AssertionError(f"API {format} error: {data['error']}")
                    time.sleep(.5)
                raise TimeoutError(f"API {format} exceeded 600s")
            preview=until_status("preview_ready")
            code,_,_=request(f"/api/v1/energy/exports/{jid}/start",method="POST")
            assert code==202
            ready=until_status("ready")
            code,ticket,_=request(f"/api/v1/energy/exports/{jid}/download-ticket",method="POST")
            assert code==200
            dl=urllib.request.urlopen(endpoint+ticket["download_path"],timeout=60)
            count=0
            while True:
                chunk=dl.read(65536)
                if not chunk: break
                count+=len(chunk)
            dl.close()
            assert count==ready["artifact"]["bytes"]
            records.append({"format":format,"elapsed_seconds":round(time.monotonic()-begin,2),
                            "source_events":preview["preview"]["source_events"],
                            "artifact_bytes":count})
            print("API_JOB_RESULT "+json.dumps(records[-1],sort_keys=True),flush=True)
        phase["value"]="after"
        time.sleep(30)
    finally:
        stop.set()
        if feeding.is_alive(): feeding.join(timeout=2)
        if polling.is_alive(): polling.join(timeout=2)
        deadline=time.monotonic()+30
        while time.monotonic()<deadline and state.snapshot()["persisted_total"] < len(accept):
            time.sleep(.1)
        saved=state.snapshot()
        ingestor.stop()
    with Session(database.engine) as session:
        persisted = {
            str(row.event_id): (row.captured_at, row.received_at)
            for row in session.execute(
                select(TelemetrySample.event_id, TelemetrySample.captured_at, TelemetrySample.received_at)
                .where(TelemetrySample.source == source_label)
            ).all()
        }
    database.dispose()
    phase_lag={"before":[],"during":[],"after":[]}
    for event_id,(when,kind,good) in accept.items():
        if not good or event_id not in persisted:
            errors.append(("not_persisted",event_id))
            continue
        captured,received=persisted[event_id]
        lag_ms=(received-captured).total_seconds()*1000
        phase_lag[kind].append(lag_ms)
    api_latency={kind:[ms for phase_name,ms in samples if phase_name==kind]
                 for kind in ("before","during","after")}
    baseline_api=percentile(api_latency["before"],.95)
    baseline_lag=percentile(phase_lag["before"],.95)
    budgets={"p95_api_ms":max(250,3*baseline_api),
             "p95_persistence_ms":max(250,3*baseline_lag),
             "max_persistence_lag_ms":2000}
    p95={kind:{"api_ms":round(percentile(api_latency[kind],.95),3),
               "persistence_ms":round(percentile(phase_lag[kind],.95),3)}
         for kind in ("before","during","after")}
    max_lag=max((max(v) for v in phase_lag.values() if v),default=float("inf"))
    if any(p95[kind]["api_ms"]>budgets["p95_api_ms"]
           or p95[kind]["persistence_ms"]>budgets["p95_persistence_ms"]
           for kind in ("during","after")):
        errors.append(("relative_p95_budget_failed",p95))
    if max_lag>budgets["max_persistence_lag_ms"]:
        errors.append(("max_lag_failed",max_lag))
    if saved["persisted_total"]!=len(accept) or len(persisted)!=len(accept):
        errors.append(("ingest_mismatch",len(accept),saved["persisted_total"],len(persisted)))
    if len(records)!=2 or any(row["elapsed_seconds"]>600 for row in records):
        errors.append(("jobs_outside_budget",records))
    result={"mode":"native_isolated_api","result":"PASS" if not errors else "FAIL",
            "classification":"synthetic_isolated_actual_http_api_not_production",
            "source_original_events":147840,"api_catalog_series":44,
            "export_jobs":records,"submitted":len(accept),
            "persisted":len(persisted),"ingestor_persisted_total":saved["persisted_total"],
            "http_request_count":len(samples),
            "api_latency_persistence_p95_by_phase_ms":p95,
            "max_persistence_lag_ms":round(max_lag,3),"fixed_budgets":budgets,
            "errors":errors, "duration_seconds":round(time.monotonic()-start,2)}
    print(json.dumps(result,sort_keys=True,ensure_ascii=False),flush=True)
    if errors: raise SystemExit(1)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("smoke", "full", "api"), default="smoke")
    parser.add_argument("--verify-existing", action="store_true")
    parser.add_argument("--verify-aggregates", action="store_true")
    parser.add_argument("--root", type=Path, default=Path("/app/data/telemetry-ingestion/energy-1364"))
    args = parser.parse_args()
    url = os.environ.get("DATABASE_URL", "")
    try:
        address = make_url(url)
    except Exception as exc:
        raise SystemExit("FAIL CLOSED: invalid dedicated test database URL") from exc
    if (address.drivername != "postgresql+psycopg"
            or address.host != "nexolab-energy-1364-pg-ssd"
            or not re.fullmatch(r"energy1364[a-z0-9]+", address.database or "")):
        raise SystemExit("FAIL CLOSED: only named isolated test host/database is allowed")
    if os.environ.get("NEXOLAB_ENERGY_1364_ISOLATED") != "1":
        raise SystemExit("FAIL CLOSED: isolated-container declaration required")
    allowed = Path("/app/data/telemetry-ingestion").resolve()
    if (not args.root.resolve().is_relative_to(allowed) or
            not args.root.name.startswith("energy1364-") or args.root.is_symlink()):
        raise SystemExit("FAIL CLOSED: exporter root must be dedicated and inside private test spool")
    if Path("/sys/fs/cgroup/memory.max").exists():
        cap = Path("/sys/fs/cgroup/memory.max").read_text().strip()
        if cap == "max" or int(cap) > LIMITS["service_memory"]:
            raise SystemExit("FAIL CLOSED: service cgroup memory cap must be <=512MiB")
    if args.verify_existing:
        verify_existing_xlsx(url, args.root)
        return
    if args.verify_aggregates:
        verify_existing_aggregates(url, args.root)
        return
    args.root.mkdir(parents=True, exist_ok=False)
    if args.mode == "api":
        api_benchmark(url)
        return
    t0 = time.monotonic()
    register_models()
    database = Database(url)
    database.create_schema()
    series = make_series()
    ticks = 3 if args.mode == "smoke" else 7 * 24 * 20
    start = datetime.now(UTC) - timedelta(days=7, minutes=2)
    count = seed(database, series, ticks=ticks, base=start)
    with database.engine.connect() as conn:
        independently_saved = conn.scalar(select(func.count()).select_from(TelemetrySample))
    assert count == independently_saved and count > (100000 if args.mode == "full" else 100)
    with database.engine.connect() as conn:
        reference = {row[0]: row[1] for row in conn.execute(select(TelemetrySample.event_id, TelemetrySample.value).where(
            TelemetrySample.source.like("energy-1364-synthetic%")))}
    assert len(reference) == count
    options_end = start + timedelta(days=7)
    jobs = ExportJobs(database, args.root, policy=ExportPolicy(reserve_bytes=LIMITS["disk_reserve"]))
    state = RuntimeState()
    ingestor = TelemetryIngestor(database, state, queue_maxsize=1000)
    jobs.start(); ingestor.start()
    quit_ingest = threading.Event()
    submissions = {"accepted": 0, "rejected": 0}
    latency = []
    def feed():
        index = 0
        while not quit_ingest.wait(0.1):
            event = TelemetryEvent(event_id=uuid.uuid4(), **series[0].model_dump(),
                captured_at=start + timedelta(seconds=10 + index * 0.001),
                value=218 + index * 0.0001, quality="valid", source="energy-1364-isolated-ingest")
            begin = time.monotonic()
            ok = ingestor.submit_payload(json.dumps(event.normalized_payload()).encode())
            latency.append((time.monotonic() - begin) * 1000)
            submissions["accepted" if ok else "rejected"] += 1
            index += 1
    thread = threading.Thread(target=feed, daemon=True)
    records = []
    samples = {"rss_peak": 0, "cgroup_peak": 0}
    try:
        if args.mode == "full":
            time.sleep(30)  # Fixed pre-load acquisition baseline.
        thread.start()
        for step, extension in ((0, "csv"), (0, "xlsx"), (600, "csv"), (600, "xlsx")):
            opts = ExportOptions(series=series, from_at=start, to_at=options_end,
                                 step_seconds=step, statistics=["avg", "min", "max"], format=extension)
            begin = time.monotonic()
            job = jobs.create(OWNER, opts)["job_id"]
            preview = until(jobs, job, "preview_ready", begin + LIMITS["job_seconds"], samples)
            jobs.confirm(OWNER, job)
            ready = until(jobs, job, "ready", begin + LIMITS["job_seconds"], samples)
            token = jobs.ticket(OWNER, job)["ticket"]
            path, _, _ = jobs.consume_ticket(job, token)
            assert path.is_file() and path.stat().st_size > 0
            if extension == "csv":
                seen = set()
                rows = 0
                with path.open(encoding="utf-8", newline="") as f:
                    for row in csv.DictReader(f):
                        rows += 1
                        if step == 0:
                            for name, event_id in row.items():
                                if not name.endswith(" · event_id") or event_id not in reference:
                                    continue
                                field = name.removesuffix(" · event_id")
                                expected = reference[event_id]
                                actual = row[field]
                                assert (actual == "") if expected is None else (abs(float(actual) - expected) < 1e-6), (event_id, actual, expected)
                                seen.add(event_id)
                assert rows == ready["artifact"]["output_rows"]
                if step == 0:
                    assert seen == reference.keys(), (len(seen), len(reference))
            else:
                with zipfile.ZipFile(path) as z:
                    assert z.testzip() is None
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=True, data_only=True)
                try:
                    first = next(wb.worksheets[0].iter_rows(min_row=2, max_row=2, values_only=True))
                    assert any(isinstance(item, (int, float)) and not isinstance(item, bool) for item in first)
                finally:
                    wb.close()
            assert preview["preview"]["source_events"] >= independently_saved
            records.append({"step_seconds": step, "format": extension,
                            "source_events": preview["preview"]["source_events"],
                            "output_rows": ready["artifact"]["output_rows"],
                            "artifact_bytes": path.stat().st_size,
                            "job_spool_bytes": sum(p.stat().st_size for p in (args.root / job).iterdir() if p.is_file()),
                            "elapsed_seconds": round(time.monotonic() - begin, 2)})
            print("JOB_RESULT " + json.dumps(records[-1], sort_keys=True), flush=True)
            if records[-1]["elapsed_seconds"] > LIMITS["job_seconds"]:
                raise AssertionError("Export exceeded fixed 600 second budget")
        if args.mode == "full":
            time.sleep(30)
    finally:
        quit_ingest.set()
        if thread.is_alive():
            thread.join(timeout=2)
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and state.snapshot()["persisted_total"] < submissions["accepted"]:
            time.sleep(0.2)
        after = state.snapshot()
        jobs.stop(); ingestor.stop(); database.dispose()
    report = {"classification": "native-isolated-synthetic-not-production", "mode": args.mode,
              "utc_completed": datetime.now(UTC).isoformat(), "platform": os.uname().machine,
              "source_events_seeded": count, "series": len(series), "ticks": ticks,
              "exports": records, "accepted": submissions["accepted"], "rejected": submissions["rejected"],
              "persisted": after["persisted_total"], "ingestion_failures": after.get("failed_total", 0),
              "p95_enqueue_ms": sorted(latency)[int(len(latency)*0.95)] if latency else None,
              "peak_export_rss_bytes": samples["rss_peak"], "peak_cgroup_bytes": samples["cgroup_peak"],
              "duration_seconds": round(time.monotonic()-t0, 2), "limits": LIMITS}
    failures = []
    if submissions["rejected"] or after["persisted_total"] != submissions["accepted"]:
        failures.append("lost-ingestion-events")
    if samples["rss_peak"] > LIMITS["export_rss"]:
        failures.append("export-rss-over-budget")
    if samples["cgroup_peak"] > LIMITS["service_memory"]:
        failures.append("service-memory-over-budget")
    if any(row["job_spool_bytes"] > LIMITS["spool_per_job"] for row in records):
        failures.append("per-job-spool-over-budget")
    if len(records) != 4:
        failures.append("not-all-formats-completed")
    report["result"] = "PASS" if not failures else "FAIL"
    report["failures"] = failures
    print(json.dumps(report, ensure_ascii=False, sort_keys=True), flush=True)
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
