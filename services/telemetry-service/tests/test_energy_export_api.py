import csv
import io
import time
from datetime import UTC, datetime, timedelta
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.contracts import TelemetryEvent

from app.energy_export.api import create_energy_export_router
from app.energy_export.jobs import ExportJobs, ExportPolicy
from app.security.authentication import JwtAuthenticator, VerifiedIdentityClaims
from app.security.authorization import Role
from app.security.dependencies import SecurityDependencies
from app.security.repository import SecurityRepository
from tests.live_dashboard_test_support import ORG_A, ORG_B
from tests.test_energy_export_source import source_fixture
from tests.test_energy_export_source import BASE
from tests.test_live_dashboard_api import AUDIENCE, ISSUER, SECRET, auth_headers


@pytest.fixture
def export_api(tmp_path):
    database, options = source_fixture(tmp_path)
    security = SecurityRepository(database)
    for subject in ("energy-operator", "other-operator"):
        security.provision_membership(organization_id=ORG_A, claims=VerifiedIdentityClaims(provider="test-oidc", subject=subject, email=None, display_name=subject), roles={Role.ENGINEER})
    dependencies = SecurityDependencies(security, mode="jwt", authenticator=JwtAuthenticator(public_key=SECRET, algorithm="HS256", issuer=ISSUER, audience=AUDIENCE, provider="test-oidc"), default_organization_id=ORG_A)
    jobs = ExportJobs(database, tmp_path / "exports", policy=ExportPolicy(reserve_bytes=1024 * 1024))
    app = FastAPI(); app.include_router(create_energy_export_router(database, jobs, dependencies))
    with TestClient(app) as client:
        try:
            yield client, jobs, options
        finally:
            jobs.stop(); database.dispose()


def wait_status(client, job_id, wanted, *, timeout=30):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = client.get(f"/api/v1/energy/exports/{job_id}", headers=auth_headers("energy-operator"))
        assert response.status_code == 200
        value = response.json()
        if value["status"] == wanted:
            return value
        assert value["status"] not in {"failed", "canceled", "expired"}, value
        time.sleep(0.03)
    raise AssertionError(f"Export did not become {wanted}")


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_actual_spawned_preview_confirm_download_and_owner_denial(export_api, format):
    client, jobs, options = export_api
    headers = auth_headers("energy-operator")
    assert client.get("/api/v1/energy/exports/catalog", headers={"X-Organization-ID": ORG_A}).status_code == 401
    catalog = client.get("/api/v1/energy/exports/catalog", headers=headers)
    assert catalog.status_code == 200 and catalog.json()["items"][0]["meter_label"] == "W1"
    options.format = format
    created = client.post("/api/v1/energy/exports/preview", headers=headers, json=options.model_dump(mode="json", by_alias=True))
    assert created.status_code == 202, created.text
    job_id = created.json()["job_id"]
    assert len(job_id) == 32
    for method, suffix in [("get", ""), ("post", "/start"), ("post", "/cancel"), ("post", "/download-ticket")]:
        denied = getattr(client, method)(f"/api/v1/energy/exports/{job_id}{suffix}", headers=auth_headers("other-operator"))
        assert denied.status_code == 404
    prepared = wait_status(client, job_id, "preview_ready")
    assert prepared["preview"]["source_events"] == 3
    assert prepared["preview"]["output_rows"] == 3
    assert prepared["preview"]["available_steps"] == [0, 30, 60, 300, 600, 900, 1800, 3600]
    late = TelemetryEvent(event_id=uuid4(), node_id="edge-a", equipment_id="LE01MP-200", channel_id="200-voltage", metric="electrical.voltage", unit="V", captured_at=BASE + timedelta(seconds=10), value=999, quality="valid", source="late-isolated-test")
    assert jobs.database.persist(late, late.normalized_payload())
    started = client.post(f"/api/v1/energy/exports/{job_id}/start", headers=headers)
    assert started.status_code == 202, started.text
    ready = wait_status(client, job_id, "ready")
    assert ready["artifact"]["output_rows"] == 3
    ticket = client.post(f"/api/v1/energy/exports/{job_id}/download-ticket", headers=headers)
    assert ticket.status_code == 200
    download = client.get(ticket.json()["download_path"])
    assert download.status_code == 200
    assert download.headers["cache-control"] == "no-store"
    assert len(download.content) == ready["artifact"]["bytes"]
    assert client.get(ticket.json()["download_path"]).status_code == 409
    if format == "csv":
        rows = list(csv.DictReader(io.StringIO(download.content.decode("utf-8"))))
        assert sorted(float(row["U, V"]) for row in rows) == [220, 221, 230]
    else:
        workbook = load_workbook(io.BytesIO(download.content), read_only=True, data_only=True)
        assert workbook.sheetnames == ["Data_1", "Metadata"]
        rows = list(workbook["Data_1"].values)
        index = rows[0].index("U, V")
        assert sorted(row[index] for row in rows[1:]) == [220, 221, 230]
        workbook.close()


def test_cancel_and_expiry_touch_only_owned_export_files(export_api, tmp_path):
    client, jobs, options = export_api
    protected = tmp_path / "telemetry-evidence.keep"
    protected.write_text("retained")
    created = client.post("/api/v1/energy/exports/preview", headers=auth_headers("energy-operator"), json=options.model_dump(mode="json", by_alias=True)).json()
    prepared = wait_status(client, created["job_id"], "preview_ready")
    response = client.post(f"/api/v1/energy/exports/{created['job_id']}/cancel", headers=auth_headers("energy-operator"))
    assert response.status_code == 200 and response.json()["status"] == "canceled"
    source = jobs.root / created["job_id"] / "source.sqlite"
    deadline = time.monotonic() + 5
    while source.exists() and time.monotonic() < deadline:
        time.sleep(0.02)
    assert not source.exists()
    assert protected.read_text() == "retained"

    assert client.post(f"/api/v1/energy/exports/{created['job_id']}/start", headers=auth_headers("energy-operator")).status_code == 409
    second = client.post("/api/v1/energy/exports/preview", headers=auth_headers("energy-operator"), json=options.model_dump(mode="json", by_alias=True)).json()
    wait_status(client, second["job_id"], "preview_ready")
    jobs.clock = lambda: datetime.now(UTC) + timedelta(seconds=jobs.policy.ttl_seconds + 1)
    expired = client.get(f"/api/v1/energy/exports/{second['job_id']}", headers=auth_headers("energy-operator"))
    assert expired.status_code == 200 and expired.json()["status"] == "expired"
    assert protected.read_text() == "retained"


def test_offline_local_auth_and_revocation_rechecked_at_actual_download(tmp_path):
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from app.security.local_api import create_local_auth_router
    from app.security.local_repository import LOCAL_AUTH_PROVIDER, LocalAuthRepository
    from app.security.local_service import LocalAuthService
    from app.security.passwords import hash_password
    database, options = source_fixture(tmp_path)
    security = SecurityRepository(database)
    local_repository = LocalAuthRepository(database)
    password = "Synthetic-Energy-Export-Password-47"
    local_repository.bootstrap_account(username="energy-local", password_hash=hash_password(password), email=None, display_name="Energy local test", organization_id=ORG_A, organization_slug="org-a", organization_name="Local test", roles={Role.ADMINISTRATOR})
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    private = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()).decode("utf-8")
    public = key.public_key().public_bytes(serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo).decode("utf-8")
    service = LocalAuthService(local_repository, security, private_key=private, algorithm="RS256", issuer="urn:nexolab:energy-test-local", audience="energy-test", access_token_seconds=300, refresh_token_seconds=3600, max_failed_attempts=3, lockout_seconds=300)
    dependencies = SecurityDependencies(security, mode="local", authenticator=JwtAuthenticator(public_key=public, algorithm="RS256", issuer="urn:nexolab:energy-test-local", audience="energy-test", provider=LOCAL_AUTH_PROVIDER), default_organization_id=ORG_A, local_session_validator=service.validate_access_claims)
    jobs = ExportJobs(database, tmp_path / "local-exports", policy=ExportPolicy(reserve_bytes=1024 * 1024))
    app = FastAPI(); app.include_router(create_local_auth_router(service)); app.include_router(create_energy_export_router(database, jobs, dependencies))
    try:
        with TestClient(app) as client:
            login = client.post("/api/v1/auth/local/login", json={"username": "energy-local", "password": password})
            assert login.status_code == 200
            tokens = login.json()
            headers = {"Authorization": f"Bearer {tokens['access_token']}", "X-Organization-ID": ORG_A}
            created = client.post("/api/v1/energy/exports/preview", headers=headers, json=options.model_dump(mode="json", by_alias=True))
            assert created.status_code == 202, created.text
            job_id = created.json()["job_id"]
            def wait_local(wanted):
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    response = client.get(f"/api/v1/energy/exports/{job_id}", headers=headers)
                    assert response.status_code == 200
                    if response.json()["status"] == wanted:
                        return
                    assert response.json()["status"] != "failed", response.json()
                    time.sleep(0.02)
                raise AssertionError("Local-auth export did not finish")
            wait_local("preview_ready")
            assert client.post(f"/api/v1/energy/exports/{job_id}/start", headers=headers).status_code == 202
            wait_local("ready")
            issued = client.post(f"/api/v1/energy/exports/{job_id}/download-ticket", headers=headers)
            assert issued.status_code == 200
            assert client.post("/api/v1/auth/local/logout", json={"refresh_token": tokens["refresh_token"]}).status_code == 204
            assert client.get(issued.json()["download_path"]).status_code == 401
            assert client.get(f"/api/v1/energy/exports/{job_id}", headers=headers).status_code == 401
    finally:
        jobs.stop(); database.dispose()
