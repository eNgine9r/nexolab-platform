import csv
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook

from app.energy_export.aggregation import materialize_cells
from app.energy_export.schemas import CatalogSeries, ExportOptions
from app.energy_export.writer import safe_text, write_artifact
from tests.test_energy_export_aggregation import BASE, options, staging


def write_fixture(tmp_path, *, format="xlsx", excel_rows=1_048_576, from_at=None):
    selected = options(step=0)
    if from_at is not None:
        selected = ExportOptions(**{**selected.model_dump(), "from_at": from_at, "to_at": from_at + timedelta(hours=1)})
    selected.format = format
    path, frozen = staging(tmp_path, [(0, 220), (10, 230), (20, None, "sensor_error")], selected)
    preview = materialize_cells(path, selected, frozen, guard=lambda: None, progress=lambda _: None)
    preview.update(snapshot_at=BASE.isoformat(), source_events=3)
    equipment = CatalogSeries(**selected.series[0].model_dump(), equipment_name='=SUM(1,2)\nТест', meter_label="W1", parameter_name="Напруга")
    output = tmp_path / f"result.{format}"
    result = write_artifact(path, output, selected, [equipment], preview, guard=lambda: None, progress=lambda _: None, excel_rows=excel_rows)
    return output, result


def test_xlsx_numbers_metadata_quality_and_explicit_partition(tmp_path):
    path, result = write_fixture(tmp_path, excel_rows=3)
    assert result["output_rows"] == 3 and result["data_sheets"] == 2
    book = load_workbook(path, read_only=True, data_only=False)
    assert book.sheetnames == ["Data_1", "Data_2", "Metadata"]
    rows = list(book["Data_1"].values)
    index = rows[0].index("U, V")
    assert rows[1][index] == 220 and rows[2][index] == 230
    assert book["Data_1"]["F2"].data_type == "s"
    assert rows[1][0].endswith("+03:00") and rows[1][1].endswith("Z")
    last = list(book["Data_2"].values)[1]
    assert last[index] is None
    assert "sensor_error" in last
    metadata = dict(list(book["Metadata"].values)[1:])
    assert metadata["source_events"] == 3 and metadata["timezone"] == "Europe/Kyiv"
    book.close()


def test_csv_rectangular_unicode_quoting_formula_safe_and_numeric(tmp_path):
    path, result = write_fixture(tmp_path, format="csv")
    with path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.reader(stream))
    assert len(rows) == 4 and len({len(row) for row in rows}) == 1
    assert rows[1][5] == "'=SUM(1,2)\nТест"
    assert float(rows[1][rows[0].index("U, V")]) == 220
    assert result["bytes"] == path.stat().st_size and len(result["sha256"]) == 64
    for value in ("=1+1", " +1", "@SUM(1)", "\tname", "-text"):
        assert safe_text(value).startswith("'")
