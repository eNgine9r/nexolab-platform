from __future__ import annotations

import csv
import hashlib
import time
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.energy_export.aggregation import iter_rows, timestamp
from app.energy_export.schemas import COUNTER, KYIV, METRICS, CatalogSeries, ExportError, ExportOptions, iso, utc
from app.energy_export.source import micros


EXCEL_ROWS = 1_048_576
BASE_HEADERS = ("Дата і час · Europe/Kyiv", "timestamp_utc", "bucket_end_utc", "bucket_end_local", "node_id", "Назва обладнання", "Лічильник", "equipment_id", "collision_row")
CSV_METADATA_HEADERS = ("period_from_utc", "period_to_utc", "period_from_local", "period_to_local", "step", "statistics", "timezone", "snapshot_at_utc")


class _BudgetFile:
    def __init__(self, stream, guard):
        self.stream, self.guard = stream, guard
        self.bytes_since_check = 0
        self.last_check = time.monotonic()

    def __getattr__(self, name):
        return getattr(self.stream, name)

    def write(self, data):
        self.bytes_since_check += len(data)
        if self.bytes_since_check >= 1024 * 1024 or time.monotonic() - self.last_check > 0.5:
            self.stream.flush(); self.guard()
            self.bytes_since_check = 0; self.last_check = time.monotonic()
        return self.stream.write(data)


def safe_text(value: str) -> str:
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")) else value


def local(value: str) -> str:
    return utc(value).astimezone(KYIV).isoformat(timespec="microseconds")


def _columns(options: ExportOptions):
    assignments = {}
    counters = {}
    columns = {}
    for index, item in enumerate(options.series):
        meter_metric = item.node_id, item.equipment_id, item.metric
        slot = counters.get(meter_metric, 0)
        counters[meter_metric] = slot + 1
        key = item.metric, slot
        assignments[index] = key
        if options.step_seconds == 0:
            fields = [("value", ""), ("original_value", " · original_value"), ("observed_unit", " · observed_unit"), ("quality", " · quality"), ("status", " · status"), ("event_id", " · event_id"), ("channel_id", " · channel_id"), ("source", " · source")]
        else:
            fields = ([("value", " · EΣ end reading"), ("delta_kwh", " · delta kWh"), ("delta_status", " · delta status"), ("anchor_start_utc", " · actual anchor start UTC"), ("anchor_end_utc", " · actual anchor end UTC"), ("anchor_tolerance_seconds", " · anchor tolerance s"), ("covered_seconds", " · actual covered s")]
                      if item.metric == COUNTER else [(stat, f" · {stat}") for stat in options.statistics])
            fields += [("status", " · status"), ("source_count", " · source count"), ("valid_count", " · valid count"), ("invalid_count", " · invalid count"), ("first_utc", " · first measurement UTC"), ("last_utc", " · last measurement UTC"), ("channel_id", " · channel_id")]
        columns[key] = fields
    order = sorted(columns, key=lambda key: (list(METRICS).index(key[0]), key[1]))
    headers = list(BASE_HEADERS)
    if options.step_seconds:
        headers[0] = "Початок інтервалу · Europe/Kyiv"
        headers[1] = "bucket_start_utc"
    for metric, slot in order:
        symbol, native, _ = METRICS[metric]
        unit = "°C" if native == "degC" else "PF (ratio)" if native == "ratio" else native
        name = f"{symbol}, {unit}" + (f" · канал {slot + 1}" if slot else "")
        headers.extend(name + suffix for _, suffix in columns[metric, slot])
    return assignments, columns, order, headers


def write_artifact(staging: Path, output: Path, options: ExportOptions, catalog_series: list[CatalogSeries], preview: dict, *, guard: Callable[[], None], progress: Callable[[int], None], excel_rows: int = EXCEL_ROWS) -> dict:
    if excel_rows < 2:
        raise ValueError("Excel sheets must allow a header and at least one record")
    assignments, columns, order, headers = _columns(options)
    meters = list(dict.fromkeys((item.node_id, item.equipment_id) for item in options.series))
    names = {(item.node_id, item.equipment_id): (item.equipment_name, item.meter_label) for item in catalog_series}
    metadata_values = [iso(options.from_at), iso(options.to_at), local(iso(options.from_at)), local(iso(options.to_at)), "RAW" if options.step_seconds == 0 else f"{options.step_seconds}s", ",".join(options.statistics) if options.step_seconds else "original", "Europe/Kyiv", preview["snapshot_at"]]
    metadata = {**dict(zip(CSV_METADATA_HEADERS, metadata_values)), "source_events": preview["source_events"], "output_rows": preview["output_rows"], "format": options.format,
                "range_semantics": "[from,to); counter boundary readings are labeled anchors only",
                "raw_collision_strategy": "exact timestamp and meter; deterministic event-id rank for same-series collisions",
                "resolution_method": preview["resolution_method"], "counter_method": "actual shared floor anchors; no counter average, reset/rollover or power-integration estimate",
                "missing_data": "blank numeric cells; missing/invalid/status/count columns; no fill",
                "estimated_bytes": preview["estimated_bytes"], "size_is_estimate": True}
    workbook = None
    stream = None
    sheets = 0
    total = 0
    sheet_rows = 0
    sheet_lengths = {}
    try:
        guard()
        if options.format == "csv":
            stream = output.open("w", encoding="utf-8", newline="")
            writer = csv.writer(stream, delimiter=",", lineterminator="\r\n")
            writer.writerow(headers + list(CSV_METADATA_HEADERS))
        else:
            workbook = Workbook(write_only=True)
            sheet = None
        for (at, meter_index, collision), cells in iter_rows(staging):
            row_time = timestamp(at)
            node_id, equipment_id = meters[meter_index]
            equipment_name, label = names[node_id, equipment_id]
            end = timestamp(min(micros(options.to_at), at + options.step_seconds * 1_000_000)) if options.step_seconds else None
            values = [local(row_time), row_time, end, local(end) if end else None, node_id, equipment_name, label, equipment_id, collision]
            mapped = {}
            for index, cell in cells.items():
                mapped[assignments[index]] = {**cell, "channel_id": options.series[index].channel_id}
            for key in order:
                cell = mapped.get(key, {})
                for field, _ in columns[key]:
                    values.append(cell.get(field, "missing" if field in {"quality", "status"} else None))
            values = [safe_text(value) if isinstance(value, str) else value for value in values]
            if options.format == "csv":
                writer.writerow(values + metadata_values)
            else:
                if sheet is None or sheet_rows >= excel_rows:
                    sheets += 1
                    sheet = workbook.create_sheet(f"Data_{sheets}")
                    sheet.freeze_panes = "A2"
                    for index, heading in enumerate(headers, 1):
                        sheet.column_dimensions[get_column_letter(index)].width = min(48, max(16, len(heading) + 2))
                    heading_cells = []
                    for heading in headers:
                        cell = WriteOnlyCell(sheet, value=heading)
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="16324F")
                        heading_cells.append(cell)
                    sheet.append(heading_cells)
                    sheet_rows = 1
                sheet.append(values)
                sheet_rows += 1
                sheet_lengths[sheet.title] = sheet_rows
            total += 1
            if total % 1000 == 0:
                if stream is not None:
                    stream.flush()
                guard(); progress(total)
        if total != preview["output_rows"]:
            raise ExportError("output_count_mismatch", "Кількість рядків файлу не відповідає попередньому перегляду.")
        if workbook is not None:
            if sheets == 0:
                sheet = workbook.create_sheet("Data_1"); sheet.append(headers); sheets = 1
            for sheet in workbook.worksheets:
                sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet_lengths.get(sheet.title, 1)}"
            meta_sheet = workbook.create_sheet("Metadata")
            meta_sheet.column_dimensions["A"].width = 34
            meta_sheet.column_dimensions["B"].width = 90
            meta_sheet.append(["Поле", "Значення"])
            for key, value in metadata.items():
                meta_sheet.append([key, safe_text(value) if isinstance(value, str) else value])
            for index, item in enumerate(catalog_series, 1):
                for key, value in item.model_dump().items():
                    meta_sheet.append([f"series_{index}.{key}", safe_text(value)])
            guard()
            with output.open("w+b") as destination:
                workbook.save(_BudgetFile(destination, guard))
        if stream is not None:
            stream.close(); stream = None
        guard()
        digest = hashlib.sha256()
        with output.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk); guard()
        progress(total)
        return {"output_rows": total, "bytes": output.stat().st_size, "sha256": digest.hexdigest(), "data_sheets": sheets if workbook is not None else None}
    finally:
        if stream is not None:
            stream.close()
        if workbook is not None:
            for sheet in workbook.worksheets:
                if not sheet.closed:
                    try:
                        sheet.close()
                    except (OSError, ValueError):
                        pass
            workbook.close()
